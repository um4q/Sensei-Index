# -*- coding: utf-8 -*-
"""QOL prompt Phase A.1 - automatic Electrical workbook backups + restore.
Direct mirror of test_backups_and_lock_detection.py's Phase 16.1 coverage,
scoped to electrical_data_access.py's own, separately-isolated backup
snapshots (never the Instrumentation workbook's)."""
import time

import pytest

import electrical_data_access as eda


def test_first_save_creates_a_backup(isolated_app_dir):
    tmp_path, da = isolated_app_dir
    assert eda.list_backups() == []
    eda.add_zone("K1B Well Pad")
    row = eda.find_first_blank_row("K1B Well Pad", "eht_removal")
    eda.save_row("K1B Well Pad", "eht_removal", row, {"trace_tag": "29103-EHT-0001"})
    backups = eda.list_backups()
    assert len(backups) == 1
    assert backups[0]["name"].startswith("Electrical_Inspection_Tracker.")


def test_backup_is_a_snapshot_of_the_workbook_before_the_write(isolated_app_dir):
    tmp_path, da = isolated_app_dir
    eda.add_zone("K1B Well Pad")
    row = eda.find_first_blank_row("K1B Well Pad", "eht_removal")
    eda.save_row("K1B Well Pad", "eht_removal", row, {"trace_tag": "BEFORE"})

    row2 = eda.find_first_blank_row("K1B Well Pad", "eht_removal")
    eda.save_row("K1B Well Pad", "eht_removal", row2, {"trace_tag": "AFTER"})

    backups = eda.list_backups()
    old_time = time.time() - 3600
    import os
    os.utime(backups[0]["path"], (old_time, old_time))

    row3 = eda.find_first_blank_row("K1B Well Pad", "eht_removal")
    eda.save_row("K1B Well Pad", "eht_removal", row3, {"trace_tag": "THIRD"})

    newest_backup = eda.list_backups()[0]["path"]
    import openpyxl
    wb = openpyxl.load_workbook(newest_backup, read_only=True)
    zone_cfg = eda.load_config()["zones"][0]
    ws = wb[zone_cfg["eht_removal_sheet"]]
    tags = {ws.cell(row=r, column=1).value for r in range(4, ws.max_row + 1)}
    wb.close()
    assert "THIRD" not in tags


def test_backup_not_repeated_within_the_interval(isolated_app_dir):
    tmp_path, da = isolated_app_dir
    eda.add_zone("K1B Well Pad")
    row = eda.find_first_blank_row("K1B Well Pad", "eht_removal")
    eda.save_row("K1B Well Pad", "eht_removal", row, {"trace_tag": "29103-EHT-0001"})
    assert len(eda.list_backups()) == 1

    row2 = eda.find_first_blank_row("K1B Well Pad", "eht_removal")
    eda.save_row("K1B Well Pad", "eht_removal", row2, {"trace_tag": "29103-EHT-0002"})
    assert len(eda.list_backups()) == 1  # still just one - within backup_interval_minutes


def test_backup_interval_setting_is_shared_with_instrumentation(isolated_app_dir):
    """Electrical's backups deliberately use the SAME backup_interval_minutes/
    backup_keep settings as Instrumentation - one cadence, not two."""
    tmp_path, da = isolated_app_dir
    da.set_setting("backup_interval_minutes", 0)  # "always back up"
    eda.add_zone("K1B Well Pad")
    row = eda.find_first_blank_row("K1B Well Pad", "eht_removal")
    eda.save_row("K1B Well Pad", "eht_removal", row, {"trace_tag": "29103-EHT-0001"})
    row2 = eda.find_first_blank_row("K1B Well Pad", "eht_removal")
    eda.save_row("K1B Well Pad", "eht_removal", row2, {"trace_tag": "29103-EHT-0002"})
    assert len(eda.list_backups()) == 2


def test_backup_keep_prunes_oldest(isolated_app_dir):
    tmp_path, da = isolated_app_dir
    da.set_setting("backup_interval_minutes", 0)
    da.set_setting("backup_keep", 3)
    eda.add_zone("K1B Well Pad")
    for i in range(6):
        row = eda.find_first_blank_row("K1B Well Pad", "eht_removal")
        eda.save_row("K1B Well Pad", "eht_removal", row, {"trace_tag": f"29103-EHT-{i:04d}"})
        time.sleep(0.01)
    backups = eda.list_backups()
    assert len(backups) == 3


def test_backup_now_ignores_interval(isolated_app_dir):
    tmp_path, da = isolated_app_dir
    eda.add_zone("K1B Well Pad")
    row = eda.find_first_blank_row("K1B Well Pad", "eht_removal")
    eda.save_row("K1B Well Pad", "eht_removal", row, {"trace_tag": "29103-EHT-0001"})
    assert len(eda.list_backups()) == 1
    eda.backup_now()
    assert len(eda.list_backups()) == 2


def test_backup_now_with_no_workbook_raises_cleanly(isolated_app_dir):
    tmp_path, da = isolated_app_dir
    assert not eda.ELECTRICAL_WORKBOOK_PATH.exists()
    with pytest.raises(FileNotFoundError):
        eda.backup_now()


def test_restore_backup_copies_snapshot_over_live_file_with_safety_snapshot(isolated_app_dir):
    tmp_path, da = isolated_app_dir
    eda.add_zone("K1B Well Pad")
    row = eda.find_first_blank_row("K1B Well Pad", "eht_removal")
    eda.save_row("K1B Well Pad", "eht_removal", row, {"trace_tag": "ORIGINAL"})
    old_backup = eda.list_backups()[0]["path"]

    row2 = eda.find_first_blank_row("K1B Well Pad", "eht_removal")
    da.set_setting("backup_interval_minutes", 0)
    eda.save_row("K1B Well Pad", "eht_removal", row2, {"trace_tag": "NEWER"})
    backups_before_restore = len(eda.list_backups())

    safety = eda.restore_backup(old_backup)
    assert safety is not None
    assert len(eda.list_backups()) == backups_before_restore + 1

    tags = {r["trace_tag"] for r in eda.read_index_rows("K1B Well Pad", "eht_removal")}
    assert "NEWER" not in tags


def test_restore_backup_missing_file_raises_cleanly(isolated_app_dir):
    tmp_path, da = isolated_app_dir
    with pytest.raises(FileNotFoundError):
        eda.restore_backup(tmp_path / "electrical_backups" / "nonexistent.xlsx")


def test_list_backups_empty_when_none_taken_yet(isolated_app_dir):
    tmp_path, da = isolated_app_dir
    assert eda.list_backups() == []


def test_backup_failure_never_blocks_the_actual_save(isolated_app_dir, monkeypatch, capsys):
    tmp_path, da = isolated_app_dir

    def boom():
        raise OSError("disk full")

    monkeypatch.setattr(eda, "_write_backup_snapshot", boom)
    eda.add_zone("K1B Well Pad")
    row = eda.find_first_blank_row("K1B Well Pad", "eht_removal")
    eda.save_row("K1B Well Pad", "eht_removal", row, {"trace_tag": "29103-EHT-0001"})  # must not raise

    rows = eda.read_index_rows("K1B Well Pad", "eht_removal")
    assert any(r["trace_tag"] == "29103-EHT-0001" for r in rows)
    assert "WARNING" in capsys.readouterr().out


def test_instrumentation_and_electrical_backups_never_cross_contaminate(isolated_app_dir):
    """The two domains' backups live in separate directories with
    distinctly-prefixed filenames - one save must never produce (or be
    findable via) the other domain's backup list."""
    tmp_path, da = isolated_app_dir
    da.set_setting("backup_interval_minutes", 0)

    row = da.find_first_blank_row(29103, "transmitter")
    da.save_row(29103, "transmitter", row, {"tag": "29103-PIT-0001"})

    eda.add_zone("K1B Well Pad")
    erow = eda.find_first_blank_row("K1B Well Pad", "eht_removal")
    eda.save_row("K1B Well Pad", "eht_removal", erow, {"trace_tag": "29103-EHT-0001"})

    da_backups = da.list_backups()
    eda_backups = eda.list_backups()
    assert len(da_backups) == 1
    assert len(eda_backups) == 1
    assert da_backups[0]["path"] != eda_backups[0]["path"]
    assert da.BACKUPS_DIR != eda.ELECTRICAL_BACKUPS_DIR
