#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Turns one row of the EHT & RTD Inspection Log into a filled, ready-to-print
PDF, using EHT_and_RTD_Installation_Inspection_Report_TEMPLATE.pdf as the
blank form.

BASIC USE:
    python3 export_eht_rtd_to_pdf.py Electrical_Inspection_Tracker.xlsx --sheet "EHT RTD Log <zone>" --rows 4

Requires: pip install openpyxl pypdf
"""
import argparse
import datetime
import re
import sys
from pathlib import Path

import openpyxl
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, TextStringObject, ArrayObject

from eht_rtd_schema import LOG_COLUMNS
from eht_rtd_field_map import (
    FIELD_MAP, YES_NO_CHECKBOXES, ATC_CHECKBOXES, PRE_INS_TEST_EQUIP_LOG_FIELDS,
    COMMENTS_LINE_FIELDS, CHECKBOX_ON, CHECKBOX_OFF,  # noqa: F401
)

HERE = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) \
    else Path(__file__).resolve().parent
DEFAULT_TEMPLATE = HERE / "EHT_and_RTD_Installation_Inspection_Report_TEMPLATE.pdf"
DEFAULT_OUTPUT_DIR = HERE / "output_pdfs"
HEADER_ROW = 3
FIRST_DATA_ROW = 4
SHEET_NAME = "EHT RTD Log"


def sanitize(text, fallback):
    text = (text or "").strip()
    text = re.sub(r"[^A-Za-z0-9\-_. ]+", "", text).strip()
    return text or fallback


def cell_to_str(value):
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_column_map(ws):
    label_to_col = {}
    dupe_labels = set()
    for col in range(1, ws.max_column + 1):
        label = ws.cell(row=HEADER_ROW, column=col).value
        if label:
            label = str(label).strip()
            if label in label_to_col:
                dupe_labels.add(label)
            label_to_col[label] = col
    if dupe_labels:
        print("WARNING: these column headers appear more than once in row "
              f"{HEADER_ROW} - only the last occurrence of each will be read:")
        for d in sorted(dupe_labels):
            print("   -", d)

    field_to_col = {}
    missing = []
    for field in LOG_COLUMNS:
        col = label_to_col.get(field["label"])
        if col is None:
            missing.append(field["label"])
        else:
            field_to_col[field["id"]] = col
    if missing:
        print("WARNING: these expected columns were not found in the Log sheet header "
              f"(row {HEADER_ROW}) and will be left blank on export:")
        for m in missing:
            print("   -", m)
    return field_to_col


def build_values_for_row(ws, field_to_col, row_num):
    """Returns a dict of REAL PDF field name -> value."""
    values = {}
    comments_text = ""

    for field in LOG_COLUMNS:
        fid = field["id"]
        col = field_to_col.get(fid)
        if col is None:
            continue
        raw = cell_to_str(ws.cell(row=row_num, column=col).value)
        if not raw:
            continue

        if field["ftype"] == "choice":
            allowed = field["choices"]
            match = next((c for c in allowed if c.lower() == raw.lower()), None)
            if match is None:
                print(f"   ! row {row_num}: '{raw}' is not a valid value for "
                      f"\"{field['label']}\" (expected one of {allowed}) - leaving blank.")
                continue
            raw = match

        # --- special cases -----------------------------------------------
        if fid in ATC_CHECKBOXES:
            if raw.upper() == "Y":
                spec = ATC_CHECKBOXES[fid]
                values[spec["field"]] = spec["on"]
            continue

        if fid in YES_NO_CHECKBOXES:
            if raw.upper() == "Y":
                values[YES_NO_CHECKBOXES[fid]] = CHECKBOX_ON
            continue

        if fid == "pre_ins_test_equip_log":
            for pdf_field in PRE_INS_TEST_EQUIP_LOG_FIELDS:
                values[pdf_field] = raw
            continue

        if fid == "comments":
            comments_text = raw
            continue

        pdf_field = FIELD_MAP.get(fid)
        if pdf_field is None:
            continue
        values[pdf_field] = raw

    if comments_text:
        lines = comments_text.splitlines() or [comments_text]
        for line, pdf_field in zip(lines, COMMENTS_LINE_FIELDS):
            if line.strip():
                values[pdf_field] = line.strip()

    return values


def ensure_default_resources(writer):
    """See export_to_pdf.py's identical helper - the original PDF's
    /AcroForm has no /DR (default resources) entry, which crashes pypdf's
    field-update code."""
    acro = writer._root_object["/AcroForm"]
    helv = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
        NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
    })
    helv_ref = writer._add_object(helv)
    font_dict = DictionaryObject({NameObject("/Helv"): helv_ref})
    dr = DictionaryObject({NameObject("/Font"): font_dict})
    acro[NameObject("/DR")] = dr
    if "/DA" not in acro:
        acro[NameObject("/DA")] = TextStringObject("/Helv 0 Tf 0 g")


def fill_pdf(template_path, values, out_path, flatten=False):
    """No signature-image stamping on this form either - every sign-off
    block is hand-signed only (see eht_rtd_field_map.py's UNMAPPED_NOTE) -
    pure AcroForm fill, no reportlab overlay."""
    reader = PdfReader(str(template_path))
    writer = PdfWriter()
    writer.append(reader)
    ensure_default_resources(writer)

    for page in writer.pages:
        writer.update_page_form_field_values(page, values, flatten=flatten)
    writer.set_need_appearances_writer(not flatten)

    if flatten:
        writer.remove_annotations(subtypes="/Widget")
        if "/AcroForm" in writer._root_object:
            acro = writer._root_object["/AcroForm"]
            acro[NameObject("/Fields")] = ArrayObject()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "wb") as fh:
        writer.write(fh)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("excel_path", nargs="?", default=str(HERE / "Electrical_Inspection_Tracker.xlsx"))
    ap.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="Path to the blank fillable PDF template")
    ap.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help="Folder to write filled PDFs into")
    ap.add_argument("--rows", default="", required=False,
                     help="Comma-separated Excel row numbers to export, e.g. --rows 4,7,12")
    ap.add_argument("--flatten", action="store_true", help="Flatten the filled fields into static page content "
                                                              "(no longer editable/fillable afterward)")
    ap.add_argument("--sheet", default=SHEET_NAME, help=f"Which sheet to read from (default: '{SHEET_NAME}')")
    args = ap.parse_args()

    excel_path = Path(args.excel_path)
    template_path = Path(args.template)
    output_dir = Path(args.output_dir)

    if not excel_path.exists():
        sys.exit(f"ERROR: can't find workbook: {excel_path}")
    if not template_path.exists():
        sys.exit(f"ERROR: can't find PDF template: {template_path}")
    if not args.rows:
        sys.exit("ERROR: pass --rows 4,7,... (this form has no batch export-flag column)")

    rows = sorted({int(x) for x in args.rows.split(",") if x.strip()})

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if args.sheet not in wb.sheetnames:
        sys.exit(f"ERROR: workbook has no '{args.sheet}' sheet. Available sheets: {', '.join(wb.sheetnames)}")
    ws = wb[args.sheet]

    field_to_col = load_column_map(ws)
    trace_col = field_to_col.get("trace_number")

    print(f"Exporting {len(rows)} row(s): {rows}")
    used_names = set()
    for row_num in rows:
        values = build_values_for_row(ws, field_to_col, row_num)
        trace = cell_to_str(ws.cell(row=row_num, column=trace_col).value) if trace_col else ""
        base_name = sanitize(trace, f"Row{row_num}")
        name = base_name
        n = 2
        while name in used_names:
            name = f"{base_name} ({n})"
            n += 1
        used_names.add(name)
        out_path = output_dir / f"{name}.pdf"
        fill_pdf(template_path, values, out_path, flatten=args.flatten)
        print(f"   -> {out_path}")

    print(f"Done. {len(rows)} PDF(s) written to {output_dir}/")


if __name__ == "__main__":
    main()
