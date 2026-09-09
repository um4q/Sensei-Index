#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
One-time builder for reference_material/EHT-004_scanned_samples/
EHT-004_Hand_Filled_Reference.xlsx - Phase D of SENSEI_INDEX_2.9_ULTRACODE_PROMPT.md.

Input: a JSON file holding the transcribe-eht004-forms workflow's return
value, shape {rows: [...], pageCount: int, totalUncertain: int}, where each
row is one scanned page instance with fields matching PAGE_SCHEMA from
transcribe-eht004-forms-wf_bf97655c-d80.js (trace_number, ..., checklist:
[{item_number, result, comments, result_pass2_alt?, comments_pass2_alt?}],
uncertain_fields: [...], optionally "<field>_pass2_alt" for any simple
field where the two independent read passes disagreed).

Usage:
    python3 build_eht004_reference_workbook.py rows.json OUTPUT.xlsx

This is a STANDALONE reference artifact, never read by the live app - see
this script's own header comment for why (Phase D / the user's own words:
"i am doing this to ensure whatever you scan is entirely correct and not
a misread").
"""
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# customer_name/project_name/contract_no/location are printed static text on
# every one of the 19 scanned EHT-004 instances (confirmed while building
# eht_pre_insulation_schema.py / SENSEI_INDEX_2.9_ULTRACODE_PROMPT.md's
# Phase B field list) - they don't vary per instance, so the transcription
# workflow's PAGE_SCHEMA deliberately didn't ask for them per-page. Recorded
# here as fixed values instead, exactly as printed on the source scans.
HEADER_DEFAULTS = {
    "customer_name": "COOEC Canada Ltd.",
    "project_name": "K1B Well Pad Project",
    "contract_no": "CA23007-SCM-AGC-004",
    "location": "Module Yard",
}

CHECKLIST_LABELS = [
    "EHT Receiving Report has been Completed if required",
    "EHT Tag Matches Part # and drawings",
    "EHT is Routed and Installed as per specifications",
    "RTD's Receiving Report has been Completed if required",
    "RTD's are installed as per drawing and specification",
    "EHT/RTD's Tags are correctly labelled and installed",
    "As-built drawings completed and submitted",
    "Terminations are complete with correct wire tags",
    "SR/MI Trace Continuity test passed (Record Resistance)",
    "EHT/RTD's Released for Insulation",
]

SIMPLE_FIELD_LABELS = [
    ("trace_number", "Trace #"),
    ("eht_controller_number", "EHT Controller #"),
    ("trace_part_number", "Trace Part #"),
    ("panel_number", "Panel #"),
    ("rtds_number", "RTD(s) #"),
    ("circuit_number", "Circuit #"),
    ("trace_line_number", "Trace Line #"),
    ("rev", "Rev #"),
]

EQUIP_FIELD_LABELS = [
    ("test_equip_model", "Test Equip. Model #"),
    ("test_equip_serial", "Test Equip. Serial #"),
    ("cal_due_date", "Cal. Due Date"),
    ("megger_500_reading", "SR/MI @ 500 VDC - Reading"),
    ("megger_500_result", "SR/MI @ 500 VDC - Passed/Fail"),
    ("megger_1000_reading", "SR/MI @ 1000 VDC - Reading"),
    ("megger_1000_result", "SR/MI @ 1000 VDC - Passed/Fail"),
    ("megger_2500_reading", "SR/MI @ 2500 VDC - Reading"),
    ("megger_2500_result", "SR/MI @ 2500 VDC - Passed/Fail"),
]

HEADER_LABELS = [
    ("customer_name", "Customer Name"),
    ("project_name", "Project Name"),
    ("contract_no", "Contract #"),
    ("location", "Location"),
]

SIGNOFF_LABELS = [
    ("yanda_rep_date", "Yanda Representative - Date (yy/mm/dd)"),
    ("yanda_rep_signature", "Yanda Representative - Signature"),
    ("client_rep_date", "Client Representative - Date (yy/mm/dd)"),
    ("client_rep_signature", "Client Representative - Signature"),
]

FLAG_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
HEADER_FILL = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")
HEADER_FONT = Font(bold=True)


def build_columns():
    """Returns [(column_key, header_label), ...] in sheet order. column_key
    is either a plain field id (looked up in the row dict) or a tuple
    ('checklist', n, 'result'|'comments') for the 10x2 checklist grid."""
    cols = [("source_pdf", "Source PDF"), ("source_page", "Source Page")]
    cols += HEADER_LABELS
    cols += SIMPLE_FIELD_LABELS
    for n, label in enumerate(CHECKLIST_LABELS, start=1):
        cols.append((("checklist", n, "result"), f"{n}. {label} - Result"))
        cols.append((("checklist", n, "comments"), f"{n}. {label} - Comments"))
    cols += EQUIP_FIELD_LABELS
    cols.append(("comments", "Comments"))
    cols += SIGNOFF_LABELS
    cols.append(("_flagged", "Flagged / Uncertain Fields (double-check these against the scan)"))
    cols.append(("_alt_values", "Pass-2 Alternate Reading (where the two independent passes disagreed)"))
    return cols


def checklist_item(row, n):
    for item in row.get("checklist", []):
        if item.get("item_number") == n:
            return item
    return {}


_JSON_ARTIFACT_RE = __import__("re").compile(r"^(\{\}|\[\]|null|None)$")


def _clean(value):
    """One transcription instance (page 2's cal_due_date) came back as the
    literal string '{}' instead of an empty string - an obvious structured-
    output artifact, not anything actually written on the paper form.
    Normalize a handful of such placeholder strings to blank; the field
    stays listed in uncertain_fields regardless, so it's still flagged for
    a manual look, just not displayed as if '{}' were handwritten data."""
    if isinstance(value, str) and _JSON_ARTIFACT_RE.match(value.strip()):
        return ""
    return value


def cell_value(row, key):
    if key in HEADER_DEFAULTS:
        return HEADER_DEFAULTS[key]
    if isinstance(key, tuple) and key[0] == "checklist":
        _, n, field = key
        return _clean(checklist_item(row, n).get(field, "") or "")
    if key == "yanda_rep_signature":
        return "Signed" if row.get("yanda_rep_signed") else ""
    if key == "client_rep_signature":
        return "Signed" if row.get("client_rep_signed") else ""
    if key == "_flagged":
        return "; ".join(row.get("uncertain_fields", []) or [])
    if key == "_alt_values":
        alts = []
        for k, v in row.items():
            if k.endswith("_pass2_alt") and v:
                alts.append(f"{k[:-len('_pass2_alt')]} = {v!r}")
        for item in row.get("checklist", []):
            if item.get("result_pass2_alt"):
                alts.append(f"checklist_item_{item['item_number']}_result = {item['result_pass2_alt']!r}")
            if item.get("comments_pass2_alt"):
                alts.append(f"checklist_item_{item['item_number']}_comments = {item['comments_pass2_alt']!r}")
        return "; ".join(alts)
    return _clean(row.get(key, "") or "")


def build_workbook(rows, source_note):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EHT-004 Pre-Insulation"

    columns = build_columns()
    ws.cell(row=1, column=1, value=(
        "Reference transcription of the 19 real hand-filled EHT-004 scans - NOT live tracker "
        "data, never auto-imported anywhere. Use this to manually cross-check the app's PDF "
        "template/schema against the original scans. " + source_note
    ))
    ws.cell(row=1, column=1).font = Font(italic=True, size=9)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))

    HEADER_ROW = 3
    for col, (key, label) in enumerate(columns, start=1):
        c = ws.cell(row=HEADER_ROW, column=col, value=label)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")

    flagged_col = next(i for i, (k, _) in enumerate(columns, start=1) if k == "_flagged")

    for r_i, row in enumerate(rows):
        excel_row = HEADER_ROW + 1 + r_i
        for col, (key, _label) in enumerate(columns, start=1):
            ws.cell(row=excel_row, column=col, value=cell_value(row, key))
        if row.get("uncertain_fields"):
            for col in range(1, len(columns) + 1):
                ws.cell(row=excel_row, column=col).fill = FLAG_FILL

    # reasonable column widths
    for col, (key, label) in enumerate(columns, start=1):
        width = 42 if key in ("_flagged", "_alt_values") else 22 if "checklist" in str(key) else 16
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=HEADER_ROW + 1, column=1).coordinate

    return wb


def main():
    if len(sys.argv) != 3:
        sys.exit("Usage: python3 build_eht004_reference_workbook.py rows.json OUTPUT.xlsx")
    rows_path, out_path = Path(sys.argv[1]), Path(sys.argv[2])
    data = json.loads(rows_path.read_text())
    rows = data["rows"] if isinstance(data, dict) else data

    total_flagged = sum(len(r.get("uncertain_fields", []) or []) for r in rows)
    source_note = (
        f"{len(rows)} page instances, double-independently read and reconciled; "
        f"{total_flagged} field(s) flagged as uncertain or disagreeing between the two "
        f"read passes (highlighted below) - check those against the original PDF by eye."
    )
    wb = build_workbook(rows, source_note)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"Wrote {out_path} ({len(rows)} rows, {total_flagged} flagged fields)")


if __name__ == "__main__":
    main()
