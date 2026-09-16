#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Turns one or more rows of a Gauge Log sheet into filled, ready-to-print
PDFs (one PDF per row), using Gauge_Inspection_Record_TEMPLATE.pdf as the
blank form.

BASIC USE (exports every row whose "Export to PDF (Y/N)" column is Y):
    python3 export_gauge_to_pdf.py Equipment_Inspection_Tracker.xlsx --sheet "Gauge Log 29103"

EXPORT SPECIFIC ROWS ONLY (Excel row numbers, overrides the Y/N flag):
    python3 export_gauge_to_pdf.py Equipment_Inspection_Tracker.xlsx --sheet "Gauge Log 29103" --rows 4,7,12

EXPORT EVERY FILLED-IN ROW REGARDLESS OF THE FLAG:
    python3 export_gauge_to_pdf.py Equipment_Inspection_Tracker.xlsx --sheet "Gauge Log 29103" --all

ALSO PRODUCE ONE COMBINED PDF OF EVERYTHING EXPORTED THIS RUN:
    python3 export_gauge_to_pdf.py Equipment_Inspection_Tracker.xlsx --sheet "Gauge Log 29103" --merge

FLATTEN THE OUTPUT (bakes the values into the page, no longer fillable):
    python3 export_gauge_to_pdf.py Equipment_Inspection_Tracker.xlsx --sheet "Gauge Log 29103" --flatten

Requires: pip install openpyxl pypdf reportlab
"""
import argparse
import datetime
import re
import sys
from pathlib import Path

import openpyxl
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, TextStringObject, ArrayObject, NumberObject
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.lib.pagesizes import letter
import io

from gauge_schema import LOG_COLUMNS
from gauge_field_map import FIELD_MAP

# Frozen-aware: when bundled into an .exe by PyInstaller, __file__ points
# inside a temp extraction folder that's wiped on exit - sys.executable's
# folder is the exe's real, persistent location.
HERE = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) \
    else Path(__file__).resolve().parent
DEFAULT_TEMPLATE = HERE / "Gauge_Inspection_Record_TEMPLATE.pdf"
DEFAULT_OUTPUT_DIR = HERE / "output_pdfs"
HEADER_ROW = 3          # combined workbook has a nav-button row above the header
FIRST_DATA_ROW = 4
SHEET_NAME = "Gauge Log"

# Yanda QC Representative signature stamp - placed on page 1, just above
# the "Signature:" underline in Part 5 (underline runs from roughly
# x=85-271, y=134.6 in PDF points - located visually, since this form's
# static text is vector-outlined and doesn't extract as real text/words).
SIGNATURE_IMAGE = HERE / "assets" / "yanda_qa_signature_transparent.png"
SIGNATURE_PAGE_INDEX = 0  # this form is a single page
SIGNATURE_X = 88
SIGNATURE_Y = 137
SIGNATURE_W = 80
SIGNATURE_H = 80 * (90 / 458)  # preserve the source image's aspect ratio


def sanitize(text, fallback):
    text = (text or "").strip()
    text = re.sub(r"[^A-Za-z0-9\-_. ]+", "", text).strip()
    return text or fallback


def cell_to_str(value):
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_column_map(ws):
    """Map each schema field's label -> its column index, by reading the
    header row. This means the script still works even if columns get
    reordered or new ones are inserted, as long as the header text matches."""
    label_to_col = {}
    for col in range(1, ws.max_column + 1):
        label = ws.cell(row=HEADER_ROW, column=col).value
        if label:
            label_to_col[str(label).strip()] = col

    field_to_col = {}
    missing = []
    for field in LOG_COLUMNS:
        col = label_to_col.get(field["label"])
        if col is None:
            missing.append(field["label"])
        else:
            field_to_col[field["id"]] = col
    if missing:
        print("WARNING: these expected columns were not found in the Log sheet header "
              f"(row {HEADER_ROW}) and will be left blank on export:")
        for m in missing:
            print("   -", m)
    return field_to_col


def rows_to_export(ws, field_to_col, explicit_rows, export_all):
    if explicit_rows:
        return sorted(set(explicit_rows))

    export_col = field_to_col.get("export_flag")
    tag_col = field_to_col.get("tag")
    rows = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        flag = cell_to_str(ws.cell(row=r, column=export_col).value).upper() if export_col else ""
        has_tag = tag_col and cell_to_str(ws.cell(row=r, column=tag_col).value)
        if export_all:
            if has_tag:
                rows.append(r)
        elif flag == "Y":
            rows.append(r)
    return rows


def build_values_for_row(ws, field_to_col, row_num):
    """Returns a dict of REAL PDF field name -> value, using
    gauge_field_map.py to translate from schema/Excel field ids to the
    actual field names on the original PDF. Simpler than transmitter's/
    valve's equivalent - no checkboxes, no split fields, this form's
    Remarks box is one real multiline field, not several single-line ones -
    so a plain 1-to-1 FIELD_MAP lookup is all this needs."""
    values = {}
    for field in LOG_COLUMNS:
        fid = field["id"]
        if fid == "export_flag":
            continue
        col = field_to_col.get(fid)
        if col is None:
            continue
        raw = cell_to_str(ws.cell(row=row_num, column=col).value)
        if not raw:
            continue

        if field["ftype"] == "choice":
            allowed = field["choices"]
            match = next((c for c in allowed if c.lower() == raw.lower()), None)
            if match is None:
                print(f"   ! row {row_num}: '{raw}' is not a valid value for "
                      f"\"{field['label']}\" (expected one of {allowed}) - leaving blank.")
                continue
            raw = match

        pdf_field = FIELD_MAP.get(fid)
        if pdf_field is None:
            continue  # this item has no real fillable field on the original PDF
        values[pdf_field] = raw

    return values


def stamp_signature(writer):
    """Overlays the Yanda QC Representative signature image onto page 1,
    just above the Signature: line in Part 5."""
    if not SIGNATURE_IMAGE.exists():
        print(f"   ! Signature image not found at {SIGNATURE_IMAGE} - skipping signature stamp. "
              f"Make sure the 'assets' folder is in the same directory as this script.")
        return
    buf = io.BytesIO()
    c = rl_canvas.Canvas(buf, pagesize=letter)
    c.drawImage(str(SIGNATURE_IMAGE), SIGNATURE_X, SIGNATURE_Y,
                width=SIGNATURE_W, height=SIGNATURE_H,
                mask="auto", preserveAspectRatio=True)
    c.save()
    buf.seek(0)
    overlay_reader = PdfReader(buf)
    writer.pages[SIGNATURE_PAGE_INDEX].merge_page(overlay_reader.pages[0])


def ensure_default_resources(writer):
    """The original PDF's /AcroForm has no /DR (default resources) entry,
    which crashes pypdf's field-update code, and an empty /DR silently
    breaks --flatten (it can't find a /Helv font to render the flattened
    text with). Populate a real Helvetica font resource instead."""
    acro = writer._root_object["/AcroForm"]
    helv = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
        NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
    })
    helv_ref = writer._add_object(helv)
    font_dict = DictionaryObject({NameObject("/Helv"): helv_ref})
    dr = DictionaryObject({NameObject("/Font"): font_dict})
    acro[NameObject("/DR")] = dr
    if "/DA" not in acro:
        acro[NameObject("/DA")] = TextStringObject("/Helv 0 Tf 0 g")


# The original PDF draws Text23 (the Part 4 Remarks box) as a ~13-line-tall
# ruled area, but never actually checked "Multi-line" on the field itself
# (confirmed against its /Ff flags) - so a value with real line breaks in
# it only ever showed its first line, no matter how tall the box looked.
# Setting the multiline bit (1 << 12) ourselves is a one-line fix for what
# reads as an oversight in the source template, not a workaround for
# anything export_gauge_to_pdf.py itself is doing wrong.
MULTILINE_FLAG = 1 << 12
MULTILINE_FIELDS = ("Text23",)


def fix_multiline_fields(writer):
    acro = writer._root_object["/AcroForm"]
    for field_ref in acro["/Fields"]:
        field = field_ref.get_object()
        if field.get("/T") in MULTILINE_FIELDS:
            current = int(field.get("/Ff", 0))
            field[NameObject("/Ff")] = NumberObject(current | MULTILINE_FLAG)


# Every field on the original PDF hardcodes its own /DA to a fixed 12pt
# ("/Helv 12 Tf") instead of pypdf's documented auto-size sentinel ("0 Tf") -
# confirmed directly against the template's AcroForm, not guessed. A fixed
# size means pypdf never shrinks long values to fit a narrow field; the
# text just gets silently cut off at the field's own clip rectangle
# instead (e.g. a real "SUCTION PRESSURE" service description clipping to
# "SUCTION PRI"). Rewriting the size token to 0 turns on pypdf's built-in
# shrink-to-fit (and, for the multiline Remarks box, proper word-wrap)
# with no other change to the appearance.
_DA_FONT_SIZE_RE = re.compile(r"(/\S+)\s+[\d.]+\s+Tf")


def fix_field_autosize(writer):
    acro = writer._root_object["/AcroForm"]
    for field_ref in acro["/Fields"]:
        field = field_ref.get_object()
        da = field.get("/DA")
        if da is None:
            continue
        new_da = _DA_FONT_SIZE_RE.sub(r"\1 0 Tf", str(da))
        if new_da != str(da):
            field[NameObject("/DA")] = TextStringObject(new_da)


def isolate_page_content_state(writer):
    """The template's own page content stream opens with an un-enclosed
    `cm` (no matching `q`/`Q` around it) that rescales/flips the coordinate
    system for the rest of the page and is simply never restored - fine
    for the template's own static content, which is entirely drawn inside
    that transform, but pypdf's flatten path (writer.update_page_form_field_
    values(..., flatten=True)) appends each filled field's draw commands by
    concatenating raw bytes onto the END of that same stream
    (PdfWriter._merge_content_stream_to_page), so they silently inherit
    that leftover transform too - shrunk, mirrored, and offset relative to
    the field's real /Rect. That's the "crooked" export: every field value
    piled up, rotated, near the Remarks box. Wrapping the page's existing
    content in its own balanced q/Q here means that transform is always
    popped before flatten's appended content runs, so filled fields render
    exactly where their /Rect says regardless of what the static content
    left active. A one-line, template-level PDF defect, not something this
    script does wrong - see also fix_multiline_fields()/fix_field_autosize()
    above for the same pattern (working around confirmed defects in the
    original PDF, not guessed).
    """
    for page in writer.pages:
        if "/Contents" not in page:
            continue
        contents = page.raw_get("/Contents").get_object()
        # This template's /Contents is itself an array of several stream
        # objects (pypdf/most writers concatenate array entries back-to-
        # back, same as one big stream) - page.get_contents() hides that by
        # handing back a detached, throwaway ContentStream copy, so mutating
        # it here would silently do nothing. Getting the real object via
        # raw_get() and editing the first/last pieces in place, instead of
        # replacing /Contents with one freshly combined stream, also matters
        # for PdfWriter._merge_content_stream_to_page: for an ArrayObject it
        # just appends each filled field's own small stream as one more
        # array entry (cheap), but for a single StreamObject it decodes the
        # *entire* existing content and rewrites it as a new object on every
        # single field - collapsing this template's array into one stream
        # would have made every flattened export re-embed a full copy of
        # this form's ~700KB of static content once per field.
        elements = list(contents) if isinstance(contents, ArrayObject) else None
        first = elements[0].get_object() if elements else contents
        last = elements[-1].get_object() if elements else contents
        first.set_data(b"q\n" + first.get_data())
        if last is first:
            first.set_data(first.get_data() + b"\nQ\n")
        else:
            last.set_data(last.get_data() + b"\nQ\n")


def fill_pdf(template_path, values, out_path, flatten=False, add_signature=True):
    reader = PdfReader(str(template_path))
    writer = PdfWriter()
    writer.append(reader)
    ensure_default_resources(writer)
    fix_multiline_fields(writer)
    fix_field_autosize(writer)
    isolate_page_content_state(writer)

    for page in writer.pages:
        writer.update_page_form_field_values(page, values, flatten=flatten)
    writer.set_need_appearances_writer(not flatten)

    if add_signature:
        stamp_signature(writer)

    if flatten:
        # update_page_form_field_values(..., flatten=True) only bakes the
        # appearance into the page content - it deliberately leaves the
        # interactive widgets in place (that's documented pypdf behaviour).
        # Actually removing them, plus clearing the AcroForm's field list,
        # is what makes the PDF stop being an editable form.
        writer.remove_annotations(subtypes="/Widget")
        if "/AcroForm" in writer._root_object:
            acro = writer._root_object["/AcroForm"]
            acro[NameObject("/Fields")] = ArrayObject()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "wb") as fh:
        writer.write(fh)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("excel_path", nargs="?", default=str(HERE / "Equipment_Inspection_Tracker.xlsx"),
                     help="Path to the log workbook (default: Equipment_Inspection_Tracker.xlsx next to this script)")
    ap.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="Path to the blank fillable PDF template")
    ap.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help="Folder to write filled PDFs into")
    ap.add_argument("--rows", default="", help="Comma-separated Excel row numbers to export "
                                                 "(overrides the Export to PDF Y/N flag), e.g. --rows 4,7,12")
    ap.add_argument("--all", action="store_true", help="Export every row that has a Tag filled in, "
                                                          "ignoring the Export to PDF Y/N flag")
    ap.add_argument("--merge", action="store_true", help="Also write one combined PDF of everything exported")
    ap.add_argument("--flatten", action="store_true", help="Flatten the filled fields into static page content "
                                                              "(no longer editable/fillable afterward)")
    ap.add_argument("--sheet", default=SHEET_NAME,
                     help=f"Which sheet to read from, e.g. 'Gauge Log 29103' (default: '{SHEET_NAME}')")
    ap.add_argument("--suffix", default="",
                     help='Text to append to the filename after the Tag, e.g. --suffix "DEV." '
                          'produces "29103-PG -0903 DEV..pdf". Default: no suffix, just "<tag>.pdf".')
    ap.add_argument("--no-signature", action="store_true",
                     help="Skip the Yanda QC Representative signature stamp entirely")
    args = ap.parse_args()

    excel_path = Path(args.excel_path)
    template_path = Path(args.template)
    output_dir = Path(args.output_dir)

    if not excel_path.exists():
        sys.exit(f"ERROR: can't find workbook: {excel_path}")
    if not template_path.exists():
        sys.exit(f"ERROR: can't find PDF template: {template_path}")

    explicit_rows = [int(x) for x in args.rows.split(",") if x.strip()] if args.rows else []

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if args.sheet not in wb.sheetnames:
        sys.exit(f"ERROR: workbook has no '{args.sheet}' sheet. "
                  f"Available sheets: {', '.join(wb.sheetnames)}")
    ws = wb[args.sheet]

    field_to_col = load_column_map(ws)
    rows = rows_to_export(ws, field_to_col, explicit_rows, args.all)

    if not rows:
        print("No rows to export. Mark a row's \"Export to PDF (Y/N)\" column as Y, "
              "or pass --rows 4,7,... , or pass --all.")
        return

    print(f"Exporting {len(rows)} row(s): {rows}")
    written = []
    tag_col = field_to_col.get("tag")
    used_names = set()

    for row_num in rows:
        values = build_values_for_row(ws, field_to_col, row_num)
        tag = cell_to_str(ws.cell(row=row_num, column=tag_col).value) if tag_col else ""
        tag_part = sanitize(tag, f"Row{row_num}")
        base_name = f"{tag_part} {args.suffix}" if args.suffix else tag_part
        name = base_name
        n = 2
        while name in used_names:
            name = f"{base_name} ({n})"
            n += 1
        used_names.add(name)
        out_path = output_dir / f"{name}.pdf"
        fill_pdf(template_path, values, out_path, flatten=args.flatten,
                  add_signature=not args.no_signature)
        written.append(out_path)
        print(f"   -> {out_path}")

    if args.merge and written:
        merged_writer = PdfWriter()
        for p in written:
            merged_writer.append(PdfReader(str(p)))
        merged_writer.set_need_appearances_writer(True)
        stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        merged_path = output_dir / f"Combined_Export_{stamp}.pdf"
        with open(merged_path, "wb") as fh:
            merged_writer.write(fh)
        print(f"Combined PDF -> {merged_path}")

    print(f"Done. {len(written)} PDF(s) written to {output_dir}/")


if __name__ == "__main__":
    main()
