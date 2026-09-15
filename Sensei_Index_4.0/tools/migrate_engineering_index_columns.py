# -*- coding: utf-8 -*-
"""
Adds the plate-6a "engineering index" columns (loop_id, pid_rev,
hookup_dwg, loop_dwg, jb_number, cable_number, io_address, eng_units,
cal_due_date, datasheet_rev, doc_status, open_ecn) to every Transmitter
Log / Valve Log sheet in Equipment_Inspection_Tracker.xlsx that doesn't
already have them.

Purely additive: only ever appends new columns past whatever a sheet
already has, by label (transmitter_schema.by_section("index") /
valve_schema.by_section("index")) - never touches an existing column or
any existing data. Safe to re-run; a sheet that already has a given
column is left alone.

Run once from inside the Sensei_Index_4.0 folder:
    python tools/migrate_engineering_index_columns.py
"""
import sys
from pathlib import Path

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import copy as _copy

HERE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(HERE))

import transmitter_schema
import valve_schema

HEADER_ROW = 3
WORKBOOK_PATH = HERE / "Equipment_Inspection_Tracker.xlsx"

SHEET_PREFIXES = {
    "Transmitter Log": transmitter_schema,
    "Valve Log": valve_schema,
}


def existing_labels(ws):
    labels = set()
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=HEADER_ROW, column=col).value
        if v:
            labels.add(str(v).strip())
    return labels


def migrate_sheet(ws, schema_mod):
    new_fields = [f for f in schema_mod.by_section("index") if f["label"] not in existing_labels(ws)]
    if not new_fields:
        return 0

    header_style_src = ws.cell(row=HEADER_ROW, column=1)
    start_col = ws.max_column + 1
    for i, field in enumerate(new_fields):
        col = start_col + i
        cell = ws.cell(row=HEADER_ROW, column=col)
        cell.value = field["label"]
        if header_style_src.has_style:
            cell.font = _copy.copy(header_style_src.font)
            cell.fill = _copy.copy(header_style_src.fill)
            cell.border = _copy.copy(header_style_src.border)
            cell.alignment = _copy.copy(header_style_src.alignment)
        ws.column_dimensions[get_column_letter(col)].width = 14

        if field.get("ftype") == "choice":
            dv = DataValidation(type="list", formula1=f'"{",".join(field["choices"])}"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(col)}{schema_mod_first_data_row()}:{get_column_letter(col)}{ws.max_row}")
    return len(new_fields)


def schema_mod_first_data_row():
    return 4  # FIRST_DATA_ROW, same constant every export module uses


def main():
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=False)
    total = 0
    for sheet_name in wb.sheetnames:
        for prefix, schema_mod in SHEET_PREFIXES.items():
            if sheet_name.startswith(prefix + " ") and not sheet_name.startswith("DEL "):
                added = migrate_sheet(wb[sheet_name], schema_mod)
                if added:
                    print(f"  {sheet_name}: +{added} columns")
                    total += added
    wb.save(WORKBOOK_PATH)
    print(f"\nSaved. {total} column(s) added across the workbook.")


if __name__ == "__main__":
    main()
