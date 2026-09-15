# -*- coding: utf-8 -*-
"""
Cross-references Instrumentation Master List.xlsx against
Equipment_Inspection_Tracker.xlsx and works out what needs to change for the
K1B-focused, K1A-excluded (except Drain Tank 2) rebuild.

This is the shared logic. build_v4_tracker.py (the thing you actually run)
imports this and writes the result; this module has no side effects on its
own so it's safe to import for a dry-run/report.

SCOPE DECISION (confirmed with the project owner, 2026-09-15):
  - K1A is out of scope entirely, EXCEPT Drain Tank 2 (series 29152).
  - K1B is the heavy-focus area: K1B Pad (29103), K1B Flowline (29203),
    and the 9050 SUS system all count as K1B for this rebuild.
  - Drain Tank 1 (29151), K1A Produced Emulsion, K1A Steam Generation, and
    the 29901/29902 injection wells are excluded.

INSTRUMENT-TYPE SCOPE: only rows the app can actually track get imported -
Transmitters (get a Transmitter Inspection & Test Record) and Valves (get a
Pneumatically Actuated Valve Check Record), same two equipment types
data_access.py has always supported. A "POSITIONER" or "*VALVE POSITION
TRANSMITTER*" row is not imported as its own record - it's paired to its
parent valve tag (same area + loop number) and folded into that valve's
Positioner Model field, matching how the Valve Check Record already has
positioner fields on the same row. Everything else (gauges, switches,
thermowells, downhole fiber, motor/VFD signals, analyzer elements, relief
valves, manual valves, ...) has no PDF template in this system - it is
listed out, counted, and skipped rather than silently dropped. See
OTHER_TYPE_COUNTS in the report this module produces.
"""
import json
import re
from collections import Counter, defaultdict

import openpyxl

HEADER_ROW = 3
FIRST_DATA_ROW = 4

# ---------------------------------------------------------------------------
# Master list sheets in scope, and where each one's *new* tags should land.
#   route "append_existing" - add new rows onto the bottom of the series
#       that already exists in the tracker for this exact sheet.
#   route "new_series"      - this sheet has no clean 1:1 existing series
#       (K1B Pad's tags are historically split across five different
#       series in the app - Mod 100/200, Mainline Separator, CMM, SUS
#       Cooler - and 9050 SUS was never tracked at all), so untracked tags
#       from it go into one new, clearly-named series instead of guessing
#       which existing bucket a given tag "belongs" in.
# ---------------------------------------------------------------------------
SCOPE_SHEETS = {
    "29103-K1B Pad": {"route": "new_series", "series_number": 29104, "series_name": "K1B Pad — Additional"},
    "KIB AREA 29203 FLOWLINE": {"route": "append_existing", "series_number": 29203},
    "29152 Drain 2": {"route": "append_existing", "series_number": 29152},
    "9050 SUS": {"route": "new_series", "series_number": 9050, "series_name": "SUS"},
}

# Excluded per the K1A decision - archived (hidden, not deleted) rather than
# dropped, same philosophy data_access.remove_series() already uses.
EXCLUDED_SERIES = [29151]  # Drain Tank 1

TRANSMITTER_PREFIXES = {"PIT", "TIT", "LIT", "FIT", "AIT", "AT", "FT", "PT", "TT"}
VALVE_PREFIXES = {"PV", "FV", "LV", "TV", "XV", "KV"}
POSITIONER_PREFIXES = {"PY", "FY", "XY", "KY", "LY", "TY"}

VALVE_TYPE_BY_PREFIX = {
    "FV": "FCV", "LV": "LCV", "PV": "PCV",
    "TV": "Other", "XV": "ESDV", "KV": "ESDV",
}

TRANSMITTER_TYPE_LABEL = {
    "PIT": "Pressure Indicating Transmitter", "PT": "Pressure Transmitter",
    "TIT": "Temperature Indicating Transmitter", "TT": "Temperature Transmitter",
    "LIT": "Level Transmitter", "FIT": "Flow Transmitter", "FT": "Flow Transmitter",
    "AIT": "Analyzer Indicating Transmitter", "AT": "Analyzer Transmitter",
}


def parse_tag(raw):
    """'9050-LIT  -957-A' -> ('9050', 'LIT', '957A'). None if it doesn't
    look like an AREA-FUNC-NUM[-SUFFIX] instrument tag at all (area must be
    numeric - filters out stray non-tag text that ended up in the column)."""
    t = re.sub(r"\s+", " ", str(raw or "").strip())
    if not t:
        return None
    parts = [p.strip() for p in t.split("-") if p.strip() != ""]
    if len(parts) < 3:
        return None
    area = parts[0]
    if not re.match(r"^\d+$", area):
        return None
    func = parts[1].upper()
    num = "".join(parts[2:])
    return area, func, num


def clean_tag(area, func, num):
    return f"{area}-{func}-{num}"


def classify(func, type_desc):
    td = (type_desc or "").upper()
    if td == "POSITIONER" or "VALVE POSITION" in td:
        return "positioner_accessory"
    if func in TRANSMITTER_PREFIXES and ("TRANSMITTER" in td or "TX" in td):
        return "transmitter"
    if func in VALVE_PREFIXES and "CONTROL VALVE" in td:
        return "valve"
    return "other"


def _find_col(norm_headers, *names):
    for want in names:
        for i, h in enumerate(norm_headers):
            if h == want:
                return i
    return None


def _find_pid_col(norm_headers):
    for i, h in enumerate(norm_headers):
        if h.startswith("p") and "id" in h:
            return i
    return None


def read_master_sheet_rows(ws):
    """Every row in one master-list sheet, classified. Returns a list of
    dicts (one per instrument, including 'other'/unsupported ones so the
    report can account for every single row)."""
    headers = [str(c.value or "").strip().replace("\xa0", " ") for c in next(ws.iter_rows(min_row=1, max_row=1))]
    norm = [re.sub(r"\s+", " ", h).strip().lower() for h in headers]

    tag_i = _find_col(norm, "tag number", "tag_number")
    svc_i = _find_col(norm, "service")
    type_i = _find_col(norm, "instrument type description", "instrument_type_desc")
    pid_i = _find_pid_col(norm)
    line_i = _find_col(norm, "line number", "line_number")
    model_i = _find_col(norm, "model", "model_name")
    mfr_i = _find_col(norm, "manufacturer", "manufacturer_name")

    def cell(row, i):
        v = row[i] if i is not None and i < len(row) else None
        return str(v).strip() if v not in (None, "") else ""

    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        raw_tag = row[tag_i] if tag_i is not None and tag_i < len(row) else None
        if not raw_tag:
            continue
        parsed = parse_tag(raw_tag)
        type_desc = cell(row, type_i)
        if not parsed:
            out.append({"raw_tag": raw_tag, "tag": None, "area": None, "func": None, "num": None,
                        "type_desc": type_desc, "cls": "unparseable"})
            continue
        area, func, num = parsed
        out.append({
            "raw_tag": raw_tag, "tag": clean_tag(area, func, num), "area": area, "func": func, "num": num,
            "type_desc": type_desc, "service": cell(row, svc_i), "pid": cell(row, pid_i),
            "line": cell(row, line_i), "model": cell(row, model_i), "mfr": cell(row, mfr_i),
            "cls": classify(func, type_desc),
        })
    return out


def pair_positioners(rows):
    """Attaches each positioner_accessory row onto its parent valve row
    (same area+num) as row['positioner']. Returns (valves, unmatched_accessories)."""
    valves = [r for r in rows if r["cls"] == "valve"]
    accessories = [r for r in rows if r["cls"] == "positioner_accessory"]
    by_key = defaultdict(list)
    for v in valves:
        by_key[(v["area"], v["num"])].append(v)
    unmatched = []
    for acc in accessories:
        targets = by_key.get((acc["area"], acc["num"]))
        if targets:
            targets[0]["positioner"] = acc
        else:
            unmatched.append(acc)
    return valves, unmatched


def load_column_map(ws):
    m = {}
    for col in range(1, ws.max_column + 1):
        label = ws.cell(row=HEADER_ROW, column=col).value
        if label:
            m[str(label).strip()] = col
    return m


def existing_tags(wb, series_numbers, sheet_prefix, key_label):
    """Every non-blank key-field value (Tag or Equip #) across the given
    series' sheets of one equipment type. sheet_prefix is 'Transmitter Log'
    or 'Valve Log'."""
    tags = set()
    for num in series_numbers:
        name = f"{sheet_prefix} {num}"
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        cm = load_column_map(ws)
        key_col = cm.get(key_label)
        if not key_col:
            continue
        for r in range(FIRST_DATA_ROW, ws.max_row + 1):
            v = ws.cell(row=r, column=key_col).value
            if v not in (None, ""):
                tags.add(str(v).strip())
    return tags


def build_report(master_path, tracker_path, registry_path):
    """Runs the whole classify+pair+dedup pipeline read-only and returns a
    dict with everything build_v4_tracker.py needs to actually write, plus
    counts for the human-readable migration report.

    The "already tracked" baseline is every series CURRENTLY in
    series_registry.json (not a hardcoded list) - deliberately, so this is
    safe to re-run: a series this tool created on an earlier run (K1B Pad
    - Additional, SUS) is registered by then, so its tags count as already
    tracked and don't get appended a second time."""
    mwb = openpyxl.load_workbook(master_path, read_only=True, data_only=True)
    twb = openpyxl.load_workbook(tracker_path, data_only=True)
    with open(registry_path, encoding="utf-8") as fh:
        registered_series = [s["number"] for s in json.load(fh)["series"]]

    tracked_transmitter_tags = existing_tags(twb, registered_series, "Transmitter Log", "Tag")
    tracked_valve_tags = existing_tags(twb, registered_series, "Valve Log", "Equip #")

    per_sheet = {}
    other_type_counts = Counter()
    unmatched_accessories_total = []

    for sheet_name, route in SCOPE_SHEETS.items():
        rows = read_master_sheet_rows(mwb[sheet_name])
        valves, unmatched_acc = pair_positioners(rows)
        transmitters = [r for r in rows if r["cls"] == "transmitter"]
        unparseable = [r for r in rows if r["cls"] == "unparseable"]
        others = [r for r in rows if r["cls"] == "other"]
        for r in others:
            other_type_counts[r["type_desc"]] += 1
        unmatched_accessories_total.extend(unmatched_acc)

        new_transmitters = [r for r in transmitters if r["tag"] not in tracked_transmitter_tags]
        new_valves = [r for r in valves if r["tag"] not in tracked_valve_tags]

        per_sheet[sheet_name] = {
            "route": route,
            "total_rows": len(rows),
            "transmitters_total": len(transmitters),
            "valves_total": len(valves),
            "other_total": len(others),
            "unparseable": unparseable,
            "already_tracked_transmitters": len(transmitters) - len(new_transmitters),
            "already_tracked_valves": len(valves) - len(new_valves),
            "new_transmitters": new_transmitters,
            "new_valves": new_valves,
        }

    return {
        "per_sheet": per_sheet,
        "other_type_counts": other_type_counts,
        "unmatched_accessories": unmatched_accessories_total,
        "tracked_transmitter_tags": tracked_transmitter_tags,
        "tracked_valve_tags": tracked_valve_tags,
    }


if __name__ == "__main__":
    BASE = "/home/user/Sensei-Index/Sensei_Index_4.0"
    report = build_report(f"{BASE}/Instrumentation Master List.xlsx",
                           f"{BASE}/Equipment_Inspection_Tracker.xlsx",
                           f"{BASE}/series_registry.json")
    for sheet, info in report["per_sheet"].items():
        print(f"\n=== {sheet} ({info['route']['route']}) ===")
        print(f"  total rows: {info['total_rows']}")
        print(f"  transmitters: {info['transmitters_total']} total, "
              f"{info['already_tracked_transmitters']} already tracked, "
              f"{len(info['new_transmitters'])} NEW to add")
        print(f"  valves: {info['valves_total']} total, "
              f"{info['already_tracked_valves']} already tracked, "
              f"{len(info['new_valves'])} NEW to add")
        print(f"  other (not tracked by this tool): {info['other_total']}")
        if info["unparseable"]:
            print(f"  UNPARSEABLE tag text: {len(info['unparseable'])}")

    print(f"\nUnmatched positioner/accessory rows (no parent valve found): {len(report['unmatched_accessories'])}")
    print(f"\nOther-type breakdown (not imported, {sum(report['other_type_counts'].values())} rows total):")
    for t, n in report["other_type_counts"].most_common():
        print(f"   {n:4d}  {t}")
