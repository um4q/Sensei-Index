#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Turns one or more rows of a General Equip Install Log sheet into filled,
ready-to-print PDFs (one PDF per row), using
General_Electrical_Equipment_Installation_and_Test_Report_TEMPLATE.pdf as
the blank form. This template is 2 pages (a LibreOffice pagination
side-effect of its real source being a Word .docx) - the Torqueing Log's
last 3 rows, Comments, and the Sign-Off block are on page 2; nothing is
missing or reordered.

BASIC USE (exports every row whose "Export to PDF (Y/N)" column is Y):
    python3 export_general_equip_install_to_pdf.py Equipment_Inspection_Tracker.xlsx --sheet "General Equip Install Log 29103"

EXPORT SPECIFIC ROWS ONLY (Excel row numbers, overrides the Y/N flag):
    python3 export_general_equip_install_to_pdf.py Equipment_Inspection_Tracker.xlsx --sheet "General Equip Install Log 29103" --rows 4,7,12

EXPORT EVERY FILLED-IN ROW REGARDLESS OF THE FLAG:
    python3 export_general_equip_install_to_pdf.py Equipment_Inspection_Tracker.xlsx --sheet "General Equip Install Log 29103" --all

ALSO PRODUCE ONE COMBINED PDF OF EVERYTHING EXPORTED THIS RUN:
    python3 export_general_equip_install_to_pdf.py Equipment_Inspection_Tracker.xlsx --sheet "General Equip Install Log 29103" --merge

FLATTEN THE OUTPUT (bakes the values into the page, no longer fillable):
    python3 export_general_equip_install_to_pdf.py Equipment_Inspection_Tracker.xlsx --sheet "General Equip Install Log 29103" --flatten

Requires: pip install openpyxl pypdf reportlab
"""
import argparse
import datetime
import re
import sys
from pathlib import Path

import openpyxl
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, TextStringObject, ArrayObject, NumberObject
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.lib.pagesizes import letter
import io

from general_equip_install_schema import LOG_COLUMNS
from general_equip_install_field_map import FIELD_MAP

# Frozen-aware: when bundled into an .exe by PyInstaller, __file__ points
# inside a temp extraction folder that's wiped on exit - sys.executable's
# folder is the exe's real, persistent location.
HERE = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) \
    else Path(__file__).resolve().parent
DEFAULT_TEMPLATE = HERE / "General_Electrical_Equipment_Installation_and_Test_Report_TEMPLATE.pdf"
DEFAULT_OUTPUT_DIR = HERE / "output_pdfs"
HEADER_ROW = 3          # combined workbook has a nav-button row above the header
FIRST_DATA_ROW = 4
SHEET_NAME = "General Equip Install Log"

# Yanda Representative signature stamp - this form's sign-off block landed
# on page 2 of the 2-page source, so the stamp goes there too - located
# directly off general_equip_install_field_positions.py's own measured
# rect for the yanda_rep_signature field.
SIGNATURE_IMAGE = HERE / "assets" / "yanda_qa_signature_transparent.png"
SIGNATURE_PAGE_INDEX = 1  # page 2 (0-indexed)
SIGNATURE_X = 401
SIGNATURE_Y = 657
SIGNATURE_W = 70
SIGNATURE_H = 70 * (90 / 458)  # preserve the source image's aspect ratio


def sanitize(text, fallback):
    text = (text or "").strip()
    text = re.sub(r"[^A-Za-z0-9\-_. ]+", "", text).strip()
    return text or fallback


def cell_to_str(value):
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_column_map(ws):
    """Map each schema field's label -> its column index, by reading the
    header row. This means the script still works even if columns get
    reordered or new ones are inserted, as long as the header text matches."""
    label_to_col = {}
    for col in range(1, ws.max_column + 1):
        label = ws.cell(row=HEADER_ROW, column=col).value
        if label:
            label_to_col[str(label).strip()] = col

    field_to_col = {}
    missing = []
    for field in LOG_COLUMNS:
        col = label_to_col.get(field["label"])
        if col is None:
            missing.append(field["label"])
        else:
            field_to_col[field["id"]] = col
    if missing:
        print("WARNING: these expected columns were not found in the Log sheet header "
              f"(row {HEADER_ROW}) and will be left blank on export:")
        for m in missing:
            print("   -", m)
    return field_to_col


def rows_to_export(ws, field_to_col, explicit_rows, export_all):
    if explicit_rows:
        return sorted(set(explicit_rows))

    export_col = field_to_col.get("export_flag")
    key_col = field_to_col.get("tag_number")
    rows = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        flag = cell_to_str(ws.cell(row=r, column=export_col).value).upper() if export_col else ""
        has_key = key_col and cell_to_str(ws.cell(row=r, column=key_col).value)
        if export_all:
            if has_key:
                rows.append(r)
        elif flag == "Y":
            rows.append(r)
    return rows


def build_values_for_row(ws, field_to_col, row_num):
    """Returns a dict of REAL PDF field name -> value, using
    general_equip_install_field_map.py to translate from schema/Excel
    field ids to the actual field names on the original PDF.
    yanda_rep_signature's typed value is skipped here - it's overlaid
    with the signature image instead, in fill_pdf()."""
    values = {}
    for field in LOG_COLUMNS:
        fid = field["id"]
        if fid in ("export_flag", "yanda_rep_signature"):
            continue
        col = field_to_col.get(fid)
        if col is None:
            continue
        raw = cell_to_str(ws.cell(row=row_num, column=col).value)
        if not raw:
            continue

        if field["ftype"] == "choice":
            allowed = field["choices"]
            match = next((c for c in allowed if c.lower() == raw.lower()), None)
            if match is None:
                print(f"   ! row {row_num}: '{raw}' is not a valid value for "
                      f"\"{field['label']}\" (expected one of {allowed}) - leaving blank.")
                continue
            raw = match

        pdf_field = FIELD_MAP.get(fid)
        if pdf_field is None:
            continue  # this item has no real fillable field on the original PDF
        values[pdf_field] = raw

    return values


def stamp_signature(writer):
    """Overlays the Yanda Representative signature image onto page 2, on
    top of the yanda_rep_signature field's own cell."""
    if not SIGNATURE_IMAGE.exists():
        print(f"   ! Signature image not found at {SIGNATURE_IMAGE} - skipping signature stamp. "
              f"Make sure the 'assets' folder is in the same directory as this script.")
        return
    buf = io.BytesIO()
    c = rl_canvas.Canvas(buf, pagesize=letter)
    c.drawImage(str(SIGNATURE_IMAGE), SIGNATURE_X, SIGNATURE_Y,
                width=SIGNATURE_W, height=SIGNATURE_H,
                mask="auto", preserveAspectRatio=True)
    c.save()
    buf.seek(0)
    overlay_reader = PdfReader(buf)
    writer.pages[SIGNATURE_PAGE_INDEX].merge_page(overlay_reader.pages[0])


def ensure_default_resources(writer):
    """This template's /AcroForm has no /DR (default resources) entry,
    which crashes pypdf's field-update code, and an empty /DR silently
    breaks --flatten (it can't find a /Helv font to render the flattened
    text with) - same defect category export_gauge_to_pdf.py already
    works around for its own template. Populate a real Helvetica font
    resource instead."""
    acro = writer._root_object["/AcroForm"]
    helv = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
        NameObject("/Encoding"): NameObject("/WinAnsiEncoding"),
    })
    helv_ref = writer._add_object(helv)
    font_dict = DictionaryObject({NameObject("/Helv"): helv_ref})
    dr = DictionaryObject({NameObject("/Font"): font_dict})
    acro[NameObject("/DR")] = dr
    if "/DA" not in acro:
        acro[NameObject("/DA")] = TextStringObject("/Helv 0 Tf 0 g")


MULTILINE_FLAG = 1 << 12
MULTILINE_FIELDS = ("comments", "equipment_resistance_testing_notes",
                     "equipment_insulation_resistance_testing_notes")


def fix_multiline_fields(writer):
    """These 3 boxes are ruled multi-line areas on the printed form, but
    the field widgets themselves don't set the /Ff multiline bit - same
    template-level oversight export_gauge_to_pdf.py already fixes for its
    own Remarks field. One-line fix, not a workaround for anything this
    script does wrong."""
    acro = writer._root_object["/AcroForm"]
    for field_ref in acro["/Fields"]:
        field = field_ref.get_object()
        if field.get("/T") in MULTILINE_FIELDS:
            current = int(field.get("/Ff", 0))
            field[NameObject("/Ff")] = NumberObject(current | MULTILINE_FLAG)


_DA_FONT_SIZE_RE = re.compile(r"(/\S+)\s+[\d.]+\s+Tf")


def fix_field_autosize(writer):
    """Rewrites each field's /DA font-size token to 0 (pypdf's auto-size
    sentinel) so a long value shrinks to fit its cell instead of clipping
    silently - same fix export_gauge_to_pdf.py applies."""
    acro = writer._root_object["/AcroForm"]
    for field_ref in acro["/Fields"]:
        field = field_ref.get_object()
        da = field.get("/DA")
        if da is None:
            continue
        new_da = _DA_FONT_SIZE_RE.sub(r"\1 0 Tf", str(da))
        if new_da != str(da):
            field[NameObject("/DA")] = TextStringObject(new_da)


def fill_pdf(template_path, values, out_path, flatten=False, add_signature=True):
    reader = PdfReader(str(template_path))
    writer = PdfWriter()
    writer.append(reader)
    ensure_default_resources(writer)
    fix_multiline_fields(writer)
    fix_field_autosize(writer)

    for page in writer.pages:
        writer.update_page_form_field_values(page, values, flatten=flatten)
    writer.set_need_appearances_writer(not flatten)

    if add_signature:
        stamp_signature(writer)

    if flatten:
        writer.remove_annotations(subtypes="/Widget")
        if "/AcroForm" in writer._root_object:
            acro = writer._root_object["/AcroForm"]
            acro[NameObject("/Fields")] = ArrayObject()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "wb") as fh:
        writer.write(fh)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("excel_path", nargs="?", default=str(HERE / "Equipment_Inspection_Tracker.xlsx"),
                     help="Path to the log workbook (default: Equipment_Inspection_Tracker.xlsx next to this script)")
    ap.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="Path to the blank fillable PDF template")
    ap.add_argument("--output-dir", default=str(DEFAULT_OUTPUT_DIR), help="Folder to write filled PDFs into")
    ap.add_argument("--rows", default="", help="Comma-separated Excel row numbers to export "
                                                 "(overrides the Export to PDF Y/N flag), e.g. --rows 4,7,12")
    ap.add_argument("--all", action="store_true", help="Export every row that has a Tag # filled in, "
                                                          "ignoring the Export to PDF Y/N flag")
    ap.add_argument("--merge", action="store_true", help="Also write one combined PDF of everything exported")
    ap.add_argument("--flatten", action="store_true", help="Flatten the filled fields into static page content "
                                                              "(no longer editable/fillable afterward)")
    ap.add_argument("--sheet", default=SHEET_NAME,
                     help=f"Which sheet to read from, e.g. 'General Equip Install Log 29103' (default: '{SHEET_NAME}')")
    ap.add_argument("--suffix", default="",
                     help='Text to append to the filename after the Tag #, e.g. --suffix "DEV." '
                          'produces "29103-EE-001 DEV..pdf". Default: no suffix, just "<tag>.pdf".')
    ap.add_argument("--no-signature", action="store_true",
                     help="Skip the Yanda Representative signature stamp entirely")
    args = ap.parse_args()

    excel_path = Path(args.excel_path)
    template_path = Path(args.template)
    output_dir = Path(args.output_dir)

    if not excel_path.exists():
        sys.exit(f"ERROR: can't find workbook: {excel_path}")
    if not template_path.exists():
        sys.exit(f"ERROR: can't find PDF template: {template_path}")

    explicit_rows = [int(x) for x in args.rows.split(",") if x.strip()] if args.rows else []

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if args.sheet not in wb.sheetnames:
        sys.exit(f"ERROR: workbook has no '{args.sheet}' sheet. "
                  f"Available sheets: {', '.join(wb.sheetnames)}")
    ws = wb[args.sheet]

    field_to_col = load_column_map(ws)
    rows = rows_to_export(ws, field_to_col, explicit_rows, args.all)

    if not rows:
        print("No rows to export. Mark a row's \"Export to PDF (Y/N)\" column as Y, "
              "or pass --rows 4,7,... , or pass --all.")
        return

    print(f"Exporting {len(rows)} row(s): {rows}")
    written = []
    key_col = field_to_col.get("tag_number")
    used_names = set()

    for row_num in rows:
        values = build_values_for_row(ws, field_to_col, row_num)
        key = cell_to_str(ws.cell(row=row_num, column=key_col).value) if key_col else ""
        key_part = sanitize(key, f"Row{row_num}")
        base_name = f"{key_part} {args.suffix}" if args.suffix else key_part
        name = base_name
        n = 2
        while name in used_names:
            name = f"{base_name} ({n})"
            n += 1
        used_names.add(name)
        out_path = output_dir / f"{name}.pdf"
        fill_pdf(template_path, values, out_path, flatten=args.flatten,
                  add_signature=not args.no_signature)
        written.append(out_path)
        print(f"   -> {out_path}")

    if args.merge and written:
        merged_writer = PdfWriter()
        for p in written:
            merged_writer.append(PdfReader(str(p)))
        merged_writer.set_need_appearances_writer(True)
        stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        merged_path = output_dir / f"Combined_Export_{stamp}.pdf"
        with open(merged_path, "wb") as fh:
            merged_writer.write(fh)
        print(f"Combined PDF -> {merged_path}")

    print(f"Done. {len(written)} PDF(s) written to {output_dir}/")


if __name__ == "__main__":
    main()
